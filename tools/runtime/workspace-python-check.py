"""Execute document round trips, including compiled Linux ARM64 extensions."""
import io
import json
import platform
from pathlib import Path
import sys
import tempfile
import zipfile

import artifact_tool_v2
import numpy as np
import pandas as pd
import openpyxl
from docx import Document
from pptx import Presentation
from PIL import Image
from reportlab.pdfgen import canvas
from pypdf import PdfReader
import pypdfium2
import pdfplumber
from lxml import etree
from cryptography.hazmat.primitives import hashes
from pydantic import TypeAdapter


def main():
    assert platform.machine() == 'aarch64'
    assert int(np.array([2, 4, 6]).sum()) == 12
    assert int(pd.DataFrame({'value': [2, 4, 6]})['value'].sum()) == 12
    assert TypeAdapter(list[int]).validate_python([2, 4, 6]) == [2, 4, 6]
    digest = hashes.Hash(hashes.SHA256())
    digest.update(b'FoldGPT')
    assert len(digest.finalize()) == 32
    assert etree.fromstring(b'<test>12</test>').text == '12'
    with tempfile.TemporaryDirectory(prefix='foldgpt-workspace-python-') as directory:
        root = Path(directory)
        book = openpyxl.Workbook()
        book.active['A1'] = 12
        book.save(root / 'check.xlsx')
        assert openpyxl.load_workbook(root / 'check.xlsx').active['A1'].value == 12
        doc = Document()
        doc.add_paragraph('FoldGPT ARM64')
        doc.save(root / 'check.docx')
        assert Document(root / 'check.docx').paragraphs[0].text == 'FoldGPT ARM64'
        slides = Presentation()
        slides.slides.add_slide(slides.slide_layouts[0]).shapes.title.text = 'FoldGPT ARM64'
        slides.save(root / 'check.pptx')
        assert Presentation(root / 'check.pptx').slides[0].shapes.title.text == 'FoldGPT ARM64'
        pdf = canvas.Canvas(str(root / 'check.pdf'))
        pdf.drawString(72, 720, 'FoldGPT ARM64 12')
        pdf.save()
        assert 'FoldGPT ARM64 12' in PdfReader(root / 'check.pdf').pages[0].extract_text()
        with pdfplumber.open(root / 'check.pdf') as pdf:
            assert 'FoldGPT' in pdf.pages[0].extract_text()
        rendered = pypdfium2.PdfDocument(root / 'check.pdf')[0].render(scale=0.2).to_pil()
        rendered.save(root / 'check.png')
        assert Image.open(root / 'check.png').width > 0
        # The Python API launches its real Node RPC implementation and shuts it
        # down at exit. Importing the Python facade alone would miss ABI faults.
        workbook = artifact_tool_v2.Workbook.create()
        sheet = workbook.worksheets.add('Check')
        sheet.get_range('A1').values = [[12]]
        artifact_tool_v2.SpreadsheetFile.export_xlsx(workbook).save(root / 'artifact.xlsx')
        assert openpyxl.load_workbook(root / 'artifact.xlsx').active['A1'].value == 12
    print(json.dumps({'passed': True, 'python': sys.version.split()[0], 'arch': platform.machine(),
                      'checks': ['numpy', 'pandas', 'pydantic', 'cryptography', 'lxml', 'xlsx',
                                 'docx', 'pptx', 'pdf', 'pdf-render', 'artifact-python-rpc']}))


if __name__ == '__main__':
    main()
