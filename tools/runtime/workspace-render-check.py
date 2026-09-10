"""Render actual Office files with LibreOffice and rasterize the resulting PDFs."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from pptx import Presentation
from pypdf import PdfReader
from PIL import Image


def main():
    with tempfile.TemporaryDirectory(prefix='foldgpt-workspace-render-') as directory:
        root = Path(directory)
        doc = Document()
        doc.add_paragraph('FoldGPT document 12')
        doc.save(root / 'document.docx')
        ppt = Presentation()
        ppt.slides.add_slide(ppt.slide_layouts[0]).shapes.title.text = 'FoldGPT presentation 12'
        ppt.save(root / 'presentation.pptx')
        bundle = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(sys.executable).parents[3]
        result = subprocess.run([str(bundle / 'dependencies/bin/fallback/libreoffice'), '--headless',
            '-env:UserInstallation=' + (root / 'profile').as_uri(), '--convert-to', 'pdf',
            '--outdir', str(root), str(root / 'document.docx'), str(root / 'presentation.pptx')],
            env=dict(os.environ, SAL_ENABLE_FILE_LOCKING='1'), capture_output=True, text=True, timeout=30)
        if result.returncode:
            raise RuntimeError('LibreOffice rendering failed, exit ' + str(result.returncode) + ': ' +
                               (result.stdout + result.stderr)[-4000:])
        for name in ('document', 'presentation'):
            pdf = root / (name + '.pdf')
            assert 'FoldGPT' in PdfReader(pdf).pages[0].extract_text()
            subprocess.run(['/usr/bin/pdftoppm', '-singlefile', '-scale-to', '128', '-png',
                            str(pdf), str(root / name)], check=True, capture_output=True, timeout=10)
            assert max(Image.open(root / (name + '.png')).size) == 128
        # Exercise the actual packaged helpers, including their runtime layout
        # and subprocess contracts, rather than only our own rendering calls.
        plugins = bundle / 'plugins/openai-primary-runtime/plugins'
        environment = dict(os.environ,
            RUNTIME_NODE=str(bundle / 'dependencies/node/bin/node'),
            RUNTIME_NODE_MODULES=str(bundle / 'dependencies/node/node_modules'),
            RUNTIME_BIN_DIR=str(bundle / 'dependencies/bin/override'))
        for script, name, image_name, extra in [
            (plugins / 'documents/skills/documents/render_docx.py', 'document', 'page-1.png', ['--dpi', '72']),
            (plugins / 'presentations/skills/presentations/container_tools/render_slides.py',
             'presentation', 'slide-1.png', ['--width', '320', '--height', '240'])]:
            extension = '.docx' if name == 'document' else '.pptx'
            destination = root / (name + '-official-helper')
            result = subprocess.run([sys.executable, str(script), str(root / (name + extension)),
                '--output_dir', str(destination), *extra], cwd=root, env=environment,
                capture_output=True, text=True, timeout=30)
            if result.returncode:
                raise RuntimeError('Packaged helper failed: ' + str(script) + ': ' +
                                   (result.stdout + result.stderr)[-6000:])
            assert Image.open(destination / image_name).width > 0
    print(json.dumps({'passed': True, 'checks': ['libreoffice-docx-pdf', 'libreoffice-pptx-pdf',
        'poppler-png', 'packaged-docx-helper', 'packaged-pptx-helper']}))


if __name__ == '__main__':
    main()
